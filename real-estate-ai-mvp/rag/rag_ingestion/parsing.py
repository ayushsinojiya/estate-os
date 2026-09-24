"""Local, deterministic parsers and strictly grounded optional interpretation.

No parser follows document links or evaluates spreadsheet formulas/macros. Native
extraction is retained in full, with locators suitable for persisted provenance.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Callable


class ParsingError(ValueError):
    """The uploaded content cannot safely be used as property knowledge."""

    def __init__(self, message, retryable=True):
        super().__init__(message)
        self.retryable = retryable


ALIASES = {
    "project": "project_name", "project name": "project_name",
    "configuration": "configuration_name", "configuration name": "configuration_name",
    "type name": "configuration_name", "listing id": "property_id",
    "property id": "property_id", "project id": "project_id",
    "rera": "rera_id", "rera number": "rera_id", "rera id": "rera_id",
    "developer": "developer_name", "developer name": "developer_name",
    "property type": "property_type", "entity type": "entity_type",
    "entity kind": "entity_type", "kind": "entity_type",
    "built up area": "built_up_area", "super built up area": "super_built_up_area",
    "carpet area": "carpet_area", "plot area": "plot_area",
    "possession": "possession_status", "possession status": "possession_status",
}


def _field(value: str) -> str:
    label = re.sub(r"[\s_-]+", " ", value.strip().lower())
    return ALIASES.get(label, label.replace(" ", "_"))


def _value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None:
        return ""
    return str(value).strip()


def _table(rows, locator, warnings):
    """Treat the first nonempty row as headers; never guess a shifted schema."""
    result = []
    headers = None
    for number, cells in rows:
        cells = [_value(cell) for cell in cells]
        if not any(cells):
            continue
        if headers is None:
            while cells and not cells[-1]:
                cells.pop()
            headers = cells
            normalized = [_field(header) for header in headers]
            if any(not key for key in normalized) or len(set(normalized)) != len(normalized):
                raise ParsingError(f"Ambiguous or empty column headers at {locator}, row {number}")
            continue
        # A repeated header can appear after a page/sheet continuation.
        if cells == headers:
            continue
        if len(cells) > len(headers) and any(cells[len(headers):]):
            warnings.append(f"Skipped {locator}, row {number}: more values than headers")
            continue
        cells = (cells + [""] * len(headers))[:len(headers)]
        values = {header: cell for header, cell in zip(headers, cells) if cell}
        row_locator = {**locator, "row": number}
        result.append({"values": values, "locator": row_locator,
                       "text": "\n".join(f"{header}: {cell}" for header, cell in values.items())})
    if not result:
        return None
    return {"locator": locator, "rows": result,
            "text": "\n\n".join(row["text"] for row in result)}


def parse_file(path: Path, filename: str) -> dict:
    """Parse physically uploaded bytes; imports stay lazy for optional formats."""
    path = Path(path)
    extension = Path(filename).suffix.lower()
    allowed = {".pdf", ".xlsx", ".xls", ".csv", ".docx", ".txt", ".md"}
    if extension == ".doc":
        raise ParsingError("Legacy DOC is unsupported; save the document as DOCX or PDF", retryable=False)
    if extension not in allowed:
        raise ParsingError(f"Unsupported knowledge file format: {extension or '(missing extension)'}", retryable=False)
    if path.stat().st_size > 50 * 1024 * 1024:
        raise ParsingError("File exceeds the 50 MB limit", retryable=False)
    blocks, warnings = [], []
    structured = extension in {".xlsx", ".xls", ".csv"}
    try:
        # Encrypted OOXML is an OLE container, not an ordinary ZIP document.
        if extension in {".xlsx", ".docx"}:
            with path.open("rb") as stream:
                if stream.read(8) == bytes.fromhex("D0CF11E0A1B11AE1"):
                    raise ParsingError("Encrypted/password-protected Office documents are unsupported", retryable=False)
            with zipfile.ZipFile(path) as archive:
                if any(info.flag_bits & 1 for info in archive.infolist()):
                    raise ParsingError("Encrypted/password-protected documents are unsupported", retryable=False)
        if extension in {".csv", ".txt", ".md"}:
            try:
                text = path.read_text(encoding="utf-8-sig")
            except UnicodeDecodeError as exc:
                raise ParsingError("Text/CSV must use UTF-8 encoding; re-save the file as UTF-8") from exc
            if extension == ".csv":
                try:
                    dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                block = _table(enumerate(csv.reader(io.StringIO(text), dialect, strict=True), 1),
                               {"format": "csv"}, warnings)
                if block:
                    blocks.append(block)
            elif text.strip():
                blocks.append({"text": text, "locator": {"block": 1, "format": extension[1:]}})
        elif extension == ".xlsx":
            import openpyxl
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
            formula_workbook = None
            try:
                formula_workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
                for sheet in workbook:
                    for cached_row, source_row in zip(sheet.iter_rows(), formula_workbook[sheet.title].iter_rows()):
                        for cached_cell, source_cell in zip(cached_row, source_row):
                            if source_cell.data_type == "f" and cached_cell.value is None:
                                warnings.append(f"Omitted formula without a cached value at {sheet.title}!{source_cell.coordinate}; recalculate and save the spreadsheet")
                    try:
                        block = _table(enumerate(sheet.iter_rows(values_only=True), 1),
                                       {"sheet": sheet.title, "format": "xlsx"}, warnings)
                    except ParsingError as exc:
                        warnings.append(f"Skipped sheet {sheet.title}: {exc}")
                        continue
                    if block:
                        blocks.append(block)
            finally:
                workbook.close()
                if formula_workbook is not None:
                    formula_workbook.close()
        elif extension == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(str(path), on_demand=True)
            try:
                for sheet in workbook.sheets():
                    rows = []
                    for index in range(sheet.nrows):
                        values = []
                        for cell in sheet.row(index):
                            values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                                          if cell.ctype == xlrd.XL_CELL_DATE else cell.value)
                        rows.append((index + 1, values))
                    try:
                        block = _table(rows, {"sheet": sheet.name, "format": "xls"}, warnings)
                    except ParsingError as exc:
                        warnings.append(f"Skipped sheet {sheet.name}: {exc}")
                        continue
                    if block:
                        blocks.append(block)
            finally:
                workbook.release_resources()
        elif extension == ".docx":
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            document = Document(path)
            for number, element in enumerate(document.element.body, 1):
                locator = {"block": number, "format": "docx"}
                if element.tag.endswith("}p"):
                    text = Paragraph(element, document).text
                    if text.strip():
                        blocks.append({"text": text, "locator": locator})
                elif element.tag.endswith("}tbl"):
                    table = Table(element, document)
                    prior_warnings = len(warnings)
                    block = _table(enumerate(([cell.text for cell in row.cells] for row in table.rows), 1),
                                   locator, warnings)
                    if len(warnings) != prior_warnings:
                        raise ParsingError("DOCX table structure is ambiguous; document was not partially accepted")
                    if block:
                        blocks.append(block)
                # OCR only embedded image parts, never rasterize native text.
                # Linked images are not fetched. Keep text/locators, not image bytes.
                from docx.oxml.ns import qn
                from PIL import Image
                for image_number, blip in enumerate(element.xpath('.//a:blip'), 1):
                    relationship = blip.get(qn('r:embed'))
                    image_locator = {**locator, 'image': image_number}
                    if not relationship:
                        warnings.append(f"Skipped external image at {image_locator}; links are not fetched")
                        continue
                    with Image.open(io.BytesIO(document.part.related_parts[relationship].blob)) as image:
                        _image_block(image, image_locator, blocks, warnings)
        elif extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ParsingError("Encrypted/password-protected PDF documents are unsupported", retryable=False)
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                method = "digital_pdf"
                # A page number or short watermark is not meaningful digital
                # content for an image-bearing scanned page. This routes OCR;
                # it is not an extraction-confidence or acceptance threshold.
                has_text = any(character.isalnum() for character in text)
                sparse_text = sum(character.isalnum() for character in text) < 40
                images = page.images
                has_images = bool(images)
                if not has_text or (sparse_text and has_images):
                    # Blank pages are retained without OCR unless they contain an image.
                    if not has_images:
                        blocks.append({"text": "", "locator": {"page": number}, "method": "digital_pdf"})
                        continue
                    ocr_text = _ocr_page(path, number - 1)
                    method = "local_ocr"
                    if not ocr_text.strip():
                        raise ParsingError(f"Scanned PDF page {number} could not be read reliably")
                    text = text + "\n" + ocr_text if text.strip() else ocr_text
                blocks.append({"text": text, "locator": {"page": number}, "method": method})
                if method == 'digital_pdf':
                    for image_number, embedded in enumerate(images, 1):
                        _image_block(embedded.image, {'page': number, 'image': image_number}, blocks, warnings)
    except ParsingError:
        raise
    except ImportError as exc:
        raise ParsingError(f"Parser dependency unavailable for {extension}: {exc.name}") from exc
    except Exception as exc:
        message = str(exc)
        if "encrypt" in message.lower() or "password" in message.lower():
            raise ParsingError("Encrypted/password-protected documents are unsupported", retryable=False) from exc
        raise ParsingError(f"Could not parse {extension} file: {message}") from exc
    if not any(block["text"].strip() for block in blocks):
        raise ParsingError("File contains no usable text or records")
    return {"blocks": blocks, "structured": structured, "warnings": warnings}


def _image_block(image, locator, blocks, warnings):
    import os
    import pytesseract
    try:
        text = pytesseract.image_to_string(image, lang=os.getenv('RAG_OCR_LANGUAGES', 'eng'))
    except Exception as exc:
        raise ParsingError(f'Local OCR failed on image at {locator}: {exc}') from exc
    if text.strip():
        blocks.append({'text': text, 'locator': locator, 'method': 'local_ocr'})
    else:
        warnings.append(f'No readable text in image at {locator}; no visual characteristics inferred')


def _ocr_page(path: Path, page_number: int) -> str:
    import os
    import pypdfium2
    import pytesseract
    document = pypdfium2.PdfDocument(str(path))
    try:
        page = document[page_number]
        try:
            bitmap = page.render(scale=3)
            try:
                return pytesseract.image_to_string(bitmap.to_pil(),
                                                    lang=os.getenv("RAG_OCR_LANGUAGES", "eng"))
            finally:
                bitmap.close()
        finally:
            page.close()
    except Exception as exc:
        raise ParsingError(f"Local OCR failed on PDF page {page_number + 1}: {exc}") from exc
    finally:
        document.close()


IDENTITY_FIELDS = {"name", "project_name", "project_id", "configuration_name", "property_id",
                   "configuration_id", "configuration_type", "variant_name", "unit_number",
                   "rera_id", "developer_name", "address", "locality", "city", "state", "country", "property_type",
                   "bhk", "carpet_area", "built_up_area", "super_built_up_area", "plot_area"}


def _kind(values):
    explicit = values.get("entity_type", "").upper()
    if explicit and explicit not in {"PROJECT", "CONFIGURATION", "PROPERTY"}:
        raise ParsingError(f"Unknown entity type: {explicit}")
    if explicit:
        return explicit
    if values.get("property_id") or values.get("unit_number"):
        return "PROPERTY"
    if values.get("configuration_name"):
        return "CONFIGURATION"
    if values.get("project_name") or values.get("rera_id") or values.get("project_id"):
        if any(values.get(field) for field in ("bhk", "carpet_area", "built_up_area", "super_built_up_area")):
            raise ParsingError("Project row with unit facts needs an explicit entity type or configuration name")
        return "PROJECT"
    return "PROPERTY"


def _check_identity(kind, identity):
    if kind == "PROJECT":
        valid = any(identity.get(field) for field in ("name", "project_name", "project_id", "rera_id"))
    elif kind == "CONFIGURATION":
        valid = bool(identity.get("configuration_name") and
                     any(identity.get(field) for field in ("project_name", "project_id", "rera_id")))
    else:
        valid = bool(identity.get("property_id") or identity.get("name"))
        if not valid:
            location = identity.get("address") or identity.get("locality")
            characteristics = sum(bool(identity.get(field)) for field in
                                  ("property_type", "bhk", "carpet_area", "built_up_area",
                                   "super_built_up_area", "plot_area"))
            valid = bool(location and characteristics >= 2)
    if not valid:
        raise ParsingError(f"{kind} record lacks reliable explicit identity")


def _key(kind, identity, locator):
    payload = json.dumps([kind, identity, locator], ensure_ascii=False, sort_keys=True)
    return kind.lower() + ":" + hashlib.sha256(payload.encode()).hexdigest()[:24]


UNCERTAIN_VALUES = {"?", "-", "--", "n/a", "na", "unknown", "unclear", "tbd", "to be confirmed", "not specified"}


def _record(values, locator, quotes=None, method="deterministic", warnings=None):
    for label, value in list(values.items()):
        if _value(value).casefold() in UNCERTAIN_VALUES:
            if warnings is not None:
                warnings.append(f"Omitted uncertain field {label} at {locator}: {_value(value)}")
    values = {label: value for label, value in values.items()
              if _value(value).casefold() not in UNCERTAIN_VALUES}
    normalized = {_field(key): _value(value) for key, value in values.items() if _value(value)}
    kind = _kind(normalized)
    identity = {field: value for field, value in normalized.items() if field in IDENTITY_FIELDS}
    _check_identity(kind, identity)
    facts = {}
    for original, value in values.items():
        field = _field(original)
        if not _value(value) or field == "entity_type":
            continue
        quote, fact_locator, fact_method = quotes[original] if quotes else (f"{original}: {_value(value)}", locator, method)
        facts[field] = {"value": _value(value), "evidence": {
            "locator": fact_locator, "quote": quote, "method": fact_method,
            "quality": "source_explicit"}}
    entity = {"key": _key(kind, identity, locator), "kind": kind,
              "identity": identity, "parent_key": None, "facts": facts}
    entities = []
    if kind == "CONFIGURATION" or (kind == "PROPERTY" and any(identity.get(field) for field in ('project_id', 'project_name', 'rera_id'))):
        parent_identity = {field: value for field, value in identity.items()
                           if field in {"project_name", "project_id", "rera_id", "developer_name"}}
        parent_key = _key("PROJECT", parent_identity, {})
        parent = {"key": parent_key, "kind": "PROJECT", "identity": parent_identity,
                  "parent_key": None, "facts": {field: facts[field] for field in parent_identity}}
        entity["parent_key"] = parent_key
        entities.append(parent)
        if kind == 'PROPERTY' and identity.get('configuration_name'):
            config_identity = {field: value for field, value in identity.items()
                               if field in {'project_id', 'project_name', 'rera_id', 'configuration_id', 'configuration_name', 'configuration_type', 'variant_name'}}
            config_key = _key('CONFIGURATION', config_identity, {})
            entities.append({'key': config_key, 'kind': 'CONFIGURATION', 'identity': config_identity,
                             'parent_key': parent_key, 'facts': {field: facts[field] for field in config_identity}})
            entity['parent_key'] = config_key
    entities.append(entity)
    return entities


def _labeled_document(raw, warnings):
    values, quotes = {}, {}
    rows, unlabeled = [], []
    for block in raw["blocks"]:
        if block.get("rows"):
            for row in block["rows"]:
                rows.extend(_record(row["values"], row["locator"], warnings=warnings))
            continue
        for line in block["text"].splitlines():
            if not line.strip():
                continue
            match = re.fullmatch(r"\s*([^:=\n]{1,80})\s*[:=]\s*(\S.*)\s*", line)
            if not match:
                unlabeled.append(line)
                continue
            label, value = match.group(1).strip(), match.group(2).strip()
            same_field = next((existing for existing in values if _field(existing) == _field(label)), None)
            if same_field is not None:
                if values[same_field] != value:
                    raise ParsingError("Multiple document records require interpretation to establish boundaries")
                continue
            values[label] = value
            quotes[label] = (line, block["locator"], block.get("method", "native_text"))
    if unlabeled:
        raise ParsingError("Document contains narrative/layout content requiring configured AI interpretation")
    if values:
        rows.extend(_record(values, {"document": True}, quotes, warnings=warnings))
    return rows


def _grounded_value(value, quote):
    """AI output is source wording; normalization belongs to a later stage."""
    if isinstance(value, dict):
        return bool(value) and all(_grounded_value(item, quote) for item in value.values())
    if isinstance(value, list):
        return bool(value) and all(_grounded_value(item, quote) for item in value)
    if value is None or isinstance(value, bool):
        return False
    text = str(value).strip()
    if not text:
        return False
    return bool(re.search(r"(?<!\w)" + re.escape(text) + r"(?!\w)", quote, re.IGNORECASE))


QUALIFICATION = re.compile(r'(?<!\w)(?:no|not|without|nahi|nahin|nathi|proposed|planned|subject to approval)(?!\w)|नहीं|नाही|\u0aa8\u0aa5\u0ac0', re.IGNORECASE)


def _validate_interpretation(result, raw):
    if not isinstance(result, dict) or not isinstance(result.get("entities"), list):
        raise ParsingError("AI interpretation returned an invalid entities structure")
    evidence_blocks = []
    for block in raw["blocks"]:
        evidence_blocks.append(block)
        evidence_blocks.extend(block.get("rows", []))
    import copy
    entities = copy.deepcopy(result["entities"])
    warnings = [str(warning) for warning in result.get("warnings", [])]
    if not entities:
        raise ParsingError("AI interpretation found no usable entities")
    keys = set()
    for entity in entities:
        if not isinstance(entity, dict) or entity.get("kind") not in {"PROJECT", "CONFIGURATION", "PROPERTY"}:
            raise ParsingError("AI interpretation returned an invalid entity kind")
        key = entity.get("key")
        if not isinstance(key, str) or not key or key in keys:
            raise ParsingError("AI entity keys must be nonempty and unique within the source")
        keys.add(key)
        identity, facts = entity.get("identity"), entity.get("facts")
        if not isinstance(identity, dict) or not isinstance(facts, dict) or not facts:
            raise ParsingError("AI entity lacks explicit identity or facts")
        _check_identity(entity["kind"], identity)
        for field, fact in list(facts.items()):
            try:
                if not isinstance(fact, dict) or not isinstance(fact.get("evidence"), dict):
                    raise ParsingError(f"AI fact {field} is malformed")
                evidence = fact["evidence"]
                locator, quote = evidence.get("locator"), evidence.get("quote")
                if not isinstance(locator, dict) or not locator or not isinstance(quote, str) or not quote.strip():
                    raise ParsingError(f"AI fact {field} lacks a locator and exact source quote")
                matches = [block for block in evidence_blocks if block["locator"] == locator]
                if not matches or not any(quote in block["text"] for block in matches):
                    raise ParsingError(f"AI fact {field} cites text absent from its source locator")
                if not _grounded_value(fact.get("value"), quote):
                    raise ParsingError(f"AI fact {field} has a value not explicitly present in its quote")
                if _value(fact["value"]).casefold() in UNCERTAIN_VALUES:
                    raise ParsingError(f"AI fact {field} contains an uncertainty marker")
                # A quoted word is insufficient when the quote denies/qualifies it.
                # Conservative safeguard, not a claim of complete semantic entailment.
                if field not in identity and QUALIFICATION.search(quote) and not QUALIFICATION.search(_value(fact['value'])):
                    raise ParsingError(f"AI fact {field} comes from a negated or qualified quote")
                evidence["method"] = "ai_grounded"
            except ParsingError as exc:
                if field in identity:
                    raise
                del facts[field]
                warnings.append(f"Omitted optional field on {key}: {exc}")
        for field, value in identity.items():
            if field not in facts or facts[field]["value"] != value:
                raise ParsingError(f"AI identity field {field} must have matching grounded fact evidence")
        entity.setdefault("parent_key", None)
    by_key = {entity["key"]: entity for entity in entities}
    for entity in entities:
        parent_key = entity["parent_key"]
        if parent_key is not None:
            parent = by_key.get(parent_key)
            if not parent or entity['kind'] == 'PROJECT' or parent['kind'] not in ({'PROJECT', 'CONFIGURATION'} if entity['kind'] == 'PROPERTY' else {'PROJECT'}):
                raise ParsingError("AI returned an invalid parent relationship")
            fields = ('configuration_id', 'configuration_name') if parent['kind'] == 'CONFIGURATION' else ("project_id", "rera_id", "project_name")
            if not any(entity["identity"].get(field) and entity["identity"][field] ==
                       parent["identity"].get(field) for field in fields):
                raise ParsingError("AI parent relationship lacks explicit shared project identity")
            project_fields = ('project_id', 'rera_id', 'project_name')
            shared = [field for field in project_fields
                      if entity['identity'].get(field) and parent['identity'].get(field)]
            if any(entity['identity'][field] != parent['identity'][field] for field in shared):
                raise ParsingError('AI parent relationship has conflicting project identity')
            if parent['kind'] == 'CONFIGURATION' and not shared and not (
                entity['identity'].get('configuration_id') and
                entity['identity']['configuration_id'] == parent['identity'].get('configuration_id')
            ):
                raise ParsingError('AI named configuration relationship lacks explicit shared project identity')
    return entities, warnings


def extract_facts(raw: dict, interpreter: Callable[[dict], dict] | None = None) -> dict:
    """Deterministic first; use full-document interpretation only when needed.

    Invalid independent structured rows are skipped with warnings. Documents
    are atomic: an unresolved document fails rather than silently losing pages.
    """
    warnings = list(raw.get("warnings", []))
    entities = []
    if raw.get("structured"):
        for block in raw["blocks"]:
            for row in block.get("rows", []):
                try:
                    entities.extend(_record(row["values"], row["locator"], warnings=warnings))
                except ParsingError as exc:
                    warnings.append(f"Skipped {row['locator']}: {exc}")
        if not entities and interpreter:
            entities, ai_warnings = _validate_interpretation(interpreter(raw), raw)
            warnings.extend(ai_warnings)
    else:
        try:
            entities = _labeled_document(raw, warnings)
        except ParsingError:
            if interpreter is None:
                raise
            entities, ai_warnings = _validate_interpretation(interpreter(raw), raw)
            warnings.extend(ai_warnings)
    if not entities:
        raise ParsingError("No records with reliable explicit property/project identity were found")
    # A repeated project reference in configuration rows is one source-local parent.
    unique = {}
    for entity in entities:
        unique.setdefault(entity["key"], entity)
    return {"entities": list(unique.values()), "warnings": warnings}
