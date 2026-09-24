import copy
import shutil
from pathlib import Path

import pytest

from rag_ingestion.parsing import ParsingError, extract_facts, parse_file


def test_structured_partial_success_and_provenance(tmp_path):
    path = tmp_path / "prices.csv"
    path.write_text("property_id,price,carpet area\nP-1,92 lakh,1250 sqft\n,10 lakh,\nP-2,85 lakh,1100 sqft\n", encoding="utf-8")
    raw = parse_file(path, path.name)
    result = extract_facts(raw)
    assert len(result["entities"]) == 2
    assert len(result["warnings"]) == 1
    fact = result["entities"][0]["facts"]["price"]
    assert fact["value"] == "92 lakh"
    assert fact["evidence"]["locator"]["row"] == 2
    assert fact["evidence"]["quote"] == "price: 92 lakh"


def test_configuration_keeps_project_parent_and_variants_separate(tmp_path):
    path = tmp_path / "prices.csv"
    path.write_text("project,configuration,bhk,carpet area\nSky,Type A,3,1650 sqft\nSky,Type B,3,1650 sqft\n", encoding="utf-8")
    result = extract_facts(parse_file(path, path.name))
    parent = [item for item in result["entities"] if item["kind"] == "PROJECT"]
    children = [item for item in result["entities"] if item["kind"] == "CONFIGURATION"]
    assert len(parent) == 1
    assert len(children) == 2
    assert children[0]["key"] != children[1]["key"]
    assert all(child["parent_key"] == parent[0]["key"] for child in children)
    assert "bhk" not in parent[0]["facts"]


def test_independent_listing_does_not_require_project(tmp_path):
    path = tmp_path / "listings.csv"
    path.write_text("property type,locality,bhk,carpet area,price\nApartment,South Bopal,3,1650 sqft,92 lakh\n", encoding="utf-8")
    entities = extract_facts(parse_file(path, path.name))["entities"]
    assert len(entities) == 1
    assert entities[0]["kind"] == "PROPERTY"
    assert entities[0]["parent_key"] is None


def test_labeled_pages_retain_context_without_ai():
    raw = {"structured": False, "blocks": [
        {"text": "Project: Sky\nDeveloper: Builder", "locator": {"page": 1}},
        {"text": "Amenities: gym\nCity: Ahmedabad", "locator": {"page": 2}}]}
    result = extract_facts(raw, lambda _: pytest.fail("AI should not run"))
    entity = result["entities"][0]
    assert entity["facts"]["amenities"]["evidence"]["locator"] == {"page": 2}
    assert entity["identity"]["project_name"] == "Sky"


def test_unparsed_document_context_is_not_silently_dropped():
    raw = {"structured": False, "blocks": [
        {"text": "Project: Sky\nPrices subject to floor-specific charges.", "locator": {"page": 1}}]}
    with pytest.raises(ParsingError, match="interpretation"):
        extract_facts(raw)


def ai_fixture():
    raw = {"structured": False, "blocks": [
        {"text": "Sky has a gym.", "locator": {"page": 1}}]}
    evidence = {"locator": {"page": 1}, "quote": "Sky has a gym."}
    interpreted = {"entities": [{"key": "sky", "kind": "PROJECT", "identity": {"name": "Sky"},
        "parent_key": None, "facts": {"name": {"value": "Sky", "evidence": copy.deepcopy(evidence)},
                                     "amenities": {"value": ["gym"], "evidence": copy.deepcopy(evidence)}}}]}
    return raw, interpreted


def test_ai_receives_full_context_and_grounded_result():
    raw, interpreted = ai_fixture()
    received = []
    def interpreter(value):
        received.append(value)
        return interpreted
    result = extract_facts(raw, interpreter)
    assert received == [raw]
    assert result["entities"][0]["facts"]["amenities"]["evidence"]["method"] == "ai_grounded"


@pytest.mark.parametrize("mutation", ["quote", "locator", "value", "identity", "number", "boolean"])
def test_reject_ai_hallucination(mutation):
    raw, interpreted = ai_fixture()
    fact = interpreted["entities"][0]["facts"]["amenities"]
    if mutation == "quote":
        fact["evidence"]["quote"] = "Sky has a pool."
    elif mutation == "locator":
        fact["evidence"]["locator"] = {"page": 2}
    elif mutation == "identity":
        interpreted["entities"][0]["identity"]["name"] = "Different Project"
    elif mutation == "number":
        fact["value"] = 5
    elif mutation == "boolean":
        fact["value"] = True
    else:
        fact["value"] = ["gym", "pool"]
    if mutation == "identity":
        with pytest.raises(ParsingError):
            extract_facts(raw, lambda _: interpreted)
    else:
        result = extract_facts(raw, lambda _: interpreted)
        assert "amenities" not in result["entities"][0]["facts"]
        assert any("Omitted optional" in warning for warning in result["warnings"])


def test_unsupported_and_encrypted_files(tmp_path):
    path = tmp_path / "file.doc"
    path.write_bytes(b"content")
    with pytest.raises(ParsingError, match="Legacy DOC"):
        parse_file(path, path.name)
    path = tmp_path / "file.xlsx"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"encrypted")
    with pytest.raises(ParsingError, match="Encrypted"):
        parse_file(path, path.name)


def test_xlsx_preserves_sheets_and_dates(tmp_path):
    import datetime
    import openpyxl
    path = tmp_path / "prices.xlsx"
    book = openpyxl.Workbook()
    for sheet, identifier in ((book.active, "P-1"), (book.create_sheet("More"), "P-2")):
        sheet.append(["property id", "possession date"])
        sheet.append([identifier, datetime.datetime(2027, 1, 2)])
    book.save(path)
    raw = parse_file(path, path.name)
    assert len(raw["blocks"]) == 2
    assert raw["blocks"][1]["rows"][0]["locator"] == {"sheet": "More", "format": "xlsx", "row": 2}
    assert extract_facts(raw)["entities"][0]["facts"]["possession_date"]["value"].startswith("2027-01-02")


def test_docx_native_table_and_paragraph_order(tmp_path):
    from docx import Document
    path = tmp_path / "units.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "property id"
    table.cell(0, 1).text = "price"
    table.cell(1, 0).text = "P-1"
    table.cell(1, 1).text = "92 lakh"
    doc.save(path)
    result = extract_facts(parse_file(path, path.name))
    assert result["entities"][0]["facts"]["price"]["evidence"]["locator"]["row"] == 2


def test_encrypted_pdf_rejected(tmp_path):
    from pypdf import PdfWriter
    path = tmp_path / "secret.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("")  # Even an empty password must not bypass the policy.
    writer.write(path)
    with pytest.raises(ParsingError, match="Encrypted"):
        parse_file(path, path.name)


def test_digital_pdf_never_calls_ocr(tmp_path, monkeypatch):
    import pypdf
    from types import SimpleNamespace
    path = tmp_path / "digital.pdf"
    path.write_bytes(b"fixture")
    monkeypatch.setattr(pypdf, "PdfReader", lambda _: SimpleNamespace(is_encrypted=False, pages=[
        SimpleNamespace(extract_text=lambda: "Project: Sky", images=[])]))
    monkeypatch.setattr("rag_ingestion.parsing._ocr_page", lambda *_: pytest.fail("OCR on digital page"))
    assert extract_facts(parse_file(path, path.name))["entities"][0]["identity"]["project_name"] == "Sky"


def test_scanned_pdf_uses_local_ocr(tmp_path, monkeypatch):
    import pypdf
    from types import SimpleNamespace
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"fixture")
    monkeypatch.setattr(pypdf, "PdfReader", lambda _: SimpleNamespace(is_encrypted=False, pages=[
        SimpleNamespace(extract_text=lambda: "", images=[object()])]))
    calls = []
    def ocr(path, number):
        calls.append(number)
        return "Project: Sky"
    monkeypatch.setattr("rag_ingestion.parsing._ocr_page", ocr)
    raw = parse_file(path, path.name)
    assert calls == [0]
    assert raw["blocks"][0]["method"] == "local_ocr"


def test_uncertain_optional_fields_omitted_with_warning(tmp_path):
    path = tmp_path / "prices.csv"
    path.write_text("property id,price,facing\nP-1,TBD,East\n", encoding="utf-8")
    result = extract_facts(parse_file(path, path.name))
    assert "price" not in result["entities"][0]["facts"]
    assert "facing" in result["entities"][0]["facts"]
    assert any("uncertain field price" in warning for warning in result["warnings"])


def test_uncertain_identity_fails_record(tmp_path):
    path = tmp_path / "prices.csv"
    path.write_text("property id,price\nTBD,92 lakh\n", encoding="utf-8")
    with pytest.raises(ParsingError, match="No records"):
        extract_facts(parse_file(path, path.name))


def test_ai_core_identity_wrong_quote_is_fatal():
    raw, interpreted = ai_fixture()
    interpreted["entities"][0]["facts"]["name"]["evidence"]["quote"] = "Different name"
    with pytest.raises(ParsingError, match="absent"):
        extract_facts(raw, lambda _: interpreted)


def test_ocr_method_survives_fact_extraction():
    raw = {"structured": False, "blocks": [{"text": "Project: Sky", "locator": {"page": 1}, "method": "local_ocr"}]}
    fact = extract_facts(raw)["entities"][0]["facts"]["project_name"]
    assert fact["evidence"]["method"] == "local_ocr"


def test_bad_sheet_does_not_discard_independent_good_sheet(tmp_path):
    import openpyxl
    path = tmp_path / "sheets.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["property id", "property_id"])
    book.active.append(["P-1", "P-2"])
    sheet = book.create_sheet("Reliable")
    sheet.append(["property id", "price"])
    sheet.append(["P-3", "92 lakh"])
    book.save(path)
    result = extract_facts(parse_file(path, path.name))
    assert len(result["entities"]) == 1
    assert result["entities"][0]["identity"]["property_id"] == "P-3"
    assert any("Skipped sheet" in warning for warning in result["warnings"])


def test_formula_without_cached_result_warns_and_never_evaluates(tmp_path):
    import openpyxl
    path = tmp_path / "formula.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["property id", "price"])
    book.active.append(["P-1", "=9000000+200000"])
    book.save(path)
    result = extract_facts(parse_file(path, path.name))
    assert "price" not in result["entities"][0]["facts"]
    assert any("formula without a cached value" in warning for warning in result["warnings"])


@pytest.mark.parametrize("quote", [
    "Sky has no gym.", "Sky does not have a gym.", "Sky is without a gym.",
    "Sky gym nahi hai.", "Sky gym nahin hai.", "Sky gym नहीं है।",
    "Sky gym नाही.", "Sky gym નથી.", "Sky has a proposed gym.",
    "Sky gym is subject to approval.",
])
def test_ai_does_not_turn_negated_or_qualified_quote_into_positive_fact(quote):
    raw, interpreted = ai_fixture()
    raw["blocks"][0]["text"] = quote
    for fact in interpreted["entities"][0]["facts"].values():
        fact["evidence"]["quote"] = quote
    result = extract_facts(raw, lambda _: interpreted)
    assert "amenities" not in result["entities"][0]["facts"]
    assert result["entities"][0]["identity"]["name"] == "Sky"
    assert any("negated or qualified" in warning for warning in result["warnings"])
    assert raw["blocks"][0]["text"] == quote


def test_ai_can_preserve_explicit_negation_in_value():
    raw, interpreted = ai_fixture()
    quote = "Sky is not furnished."
    raw["blocks"][0]["text"] = quote
    facts = interpreted["entities"][0]["facts"]
    facts["name"]["evidence"]["quote"] = quote
    del facts["amenities"]
    facts["furnishing_status"] = {"value": "not furnished", "evidence": {"locator": {"page": 1}, "quote": quote}}
    assert extract_facts(raw, lambda _: interpreted)["entities"][0]["facts"]["furnishing_status"]["value"] == "not furnished"


def test_permanent_parser_rejection_has_retryable_false(tmp_path):
    path = tmp_path / "file.doc"
    path.write_bytes(b"legacy content")
    with pytest.raises(ParsingError) as error:
        parse_file(path, path.name)
    assert error.value.retryable is False
    path = tmp_path / "file.xlsx"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"encrypted")
    with pytest.raises(ParsingError) as error:
        parse_file(path, path.name)
    assert error.value.retryable is False
    assert ParsingError("OCR dependency missing").retryable is True


def test_digital_pdf_embedded_image_text_keeps_separate_provenance(tmp_path, monkeypatch):
    import pypdf
    from PIL import Image
    from types import SimpleNamespace
    path = tmp_path / 'mixed.pdf'
    path.write_bytes(b'fixture')
    monkeypatch.setattr(pypdf, 'PdfReader', lambda _: SimpleNamespace(is_encrypted=False, pages=[
        SimpleNamespace(extract_text=lambda: 'Property id: OFFICE-1\nName: Riverside Business Office',
                        images=[SimpleNamespace(image=Image.new('RGB', (300, 100)))])]))
    monkeypatch.setattr('pytesseract.image_to_string', lambda *a, **k: 'Workstations: 24')
    result = extract_facts(parse_file(path, path.name))
    fact = result['entities'][0]['facts']['workstations']
    assert fact['value'] == '24'
    assert fact['evidence']['locator'] == {'page': 1, 'image': 1}
    assert fact['evidence']['method'] == 'local_ocr'


def test_docx_embedded_image_text_is_not_dropped(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    path = tmp_path / 'office.docx'
    image = tmp_path / 'labels.png'
    Image.new('RGB', (300, 100)).save(image)
    doc = Document()
    doc.add_paragraph('Property id: OFFICE-1')
    doc.add_picture(str(image))
    doc.save(path)
    monkeypatch.setattr('pytesseract.image_to_string', lambda *a, **k: 'Workstations: 24')
    raw = parse_file(path, path.name)
    fact = extract_facts(raw)['entities'][0]['facts']['workstations']
    assert fact['value'] == '24'
    assert fact['evidence']['locator'] == {'block': 2, 'format': 'docx', 'image': 1}
    assert fact['evidence']['method'] == 'local_ocr'
    assert not any(isinstance(value, bytes) for block in raw['blocks'] for value in block.values())


def test_docx_image_without_readable_text_warns_without_inventing_facts(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    path = tmp_path / 'photo.docx'
    image = tmp_path / 'photo.png'
    Image.new('RGB', (300, 100)).save(image)
    doc = Document()
    doc.add_paragraph('Property id: OFFICE-1')
    doc.add_picture(str(image))
    doc.save(path)
    monkeypatch.setattr('pytesseract.image_to_string', lambda *a, **k: '')
    result = extract_facts(parse_file(path, path.name))
    assert set(result['entities'][0]['facts']) == {'property_id'}
    assert any('image' in warning and 'readable text' in warning for warning in result['warnings'])


@pytest.mark.skipif(shutil.which('tesseract') is None, reason='Real OCR requires installed Tesseract; included in Docker')
@pytest.mark.parametrize('format', ['docx', 'pdf'])
def test_real_local_ocr_extracts_explicit_image_labels(tmp_path, format):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new('RGB', (1000, 240), 'white')
    drawing = ImageDraw.Draw(image)
    drawing.multiline_text((30, 30), 'Property id: OFFICE-1\nWorkstations: 24',
                           font=ImageFont.load_default(size=38), fill='black', spacing=25)
    path = tmp_path / ('labels.' + format)
    if format == 'pdf':
        image.save(path, 'PDF', resolution=150)
    else:
        from docx import Document
        png = tmp_path / 'labels.png'
        image.save(png)
        document = Document()
        document.add_picture(str(png))
        document.save(path)
    result = extract_facts(parse_file(path, path.name))
    assert result['entities'][0]['identity']['property_id'] == 'OFFICE-1'
    assert result['entities'][0]['facts']['workstations']['value'] == '24'
    assert result['entities'][0]['facts']['workstations']['evidence']['method'] == 'local_ocr'


def test_ai_same_configuration_name_cannot_link_different_projects():
    text = 'Sky offers Type A. Other offers Type A listing L-1.'
    raw = {'structured': False, 'blocks': [{'text': text, 'locator': {'page': 1}}]}
    def entity(key, kind, identity, parent):
        return {'key': key, 'kind': kind, 'identity': identity, 'parent_key': parent,
                'facts': {field: {'value': value, 'evidence': {'quote': text, 'locator': {'page': 1}}}
                          for field, value in identity.items()}}
    interpreted = {'entities': [
        entity('p', 'PROJECT', {'project_name': 'Sky'}, None),
        entity('c', 'CONFIGURATION', {'project_name': 'Sky', 'configuration_name': 'Type A'}, 'p'),
        entity('u', 'PROPERTY', {'project_name': 'Other', 'configuration_name': 'Type A', 'property_id': 'L-1'}, 'c'),
    ]}
    with pytest.raises(ParsingError, match='project identity'):
        extract_facts(raw, lambda _: interpreted)
