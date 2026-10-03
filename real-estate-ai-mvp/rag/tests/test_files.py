"""Upload validation and deterministic parsing of office files."""

import io

import pytest

from rag_service.ingest.files import UnsupportedFile, detect, docx_pages, pptx_pages, workbook_pages


def test_declared_type_must_match_the_bytes():
    with pytest.raises(UnsupportedFile):
        detect("brochure.pdf", b"not really a pdf")
    with pytest.raises(UnsupportedFile):
        detect("plan.png", b"GIF89a....")
    with pytest.raises(UnsupportedFile):
        detect("malware.exe", b"MZ")
    with pytest.raises(UnsupportedFile):
        detect("empty.txt", b"")
    assert detect("b.PDF", b"%PDF-1.7 ...").kind == "pdf"
    assert detect("x.webp", b"RIFF\x00\x00\x00\x00WEBPVP8 ").mime == "image/webp"


def test_csv_becomes_a_table():
    pages = workbook_pages("csv", "Stage,Percent\nBooking,10%\nPlinth,15%\n".encode())
    assert pages[0].splitlines()[2:] == ["| Stage | Percent |", "|---|---|", "| Booking | 10% |", "| Plinth | 15% |"]


def test_docx_headings_and_tables_survive():
    import docx

    document = docx.Document()
    document.add_heading("Payment plan", level=1)
    document.add_paragraph("Construction-linked plan for Tower A.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Stage", "Due"
    table.cell(1, 0).text, table.cell(1, 1).text = "Booking", "10%"
    buffer = io.BytesIO()
    document.save(buffer)
    md = docx_pages(buffer.getvalue())[0]
    assert md.startswith("# Payment plan")
    assert "| Stage | Due |" in md and "| Booking | 10% |" in md


def test_pptx_slides_become_pages():
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Amenities"
    slide.placeholders[1].text_frame.text = "Clubhouse"
    slide.placeholders[1].text_frame.add_paragraph().text = "Swimming pool"
    buffer = io.BytesIO()
    deck.save(buffer)
    pages = pptx_pages(buffer.getvalue())
    assert pages[0].startswith("## Amenities") and "- Swimming pool" in pages[0]
